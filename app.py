"""批次文件審核輔助 MVP：所有輸出都必須人工覆核，禁止自動放行。"""
import argparse
import calendar
import hashlib
import json
import logging
import re
import shutil
import sqlite3
import tempfile
import uuid
from datetime import date
from pathlib import Path
from typing import Literal
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, HTMLResponse
from pydantic import BaseModel, Field, field_validator
from PIL import Image, ImageOps
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'output'; OUT.mkdir(exist_ok=True)
logging.basicConfig(level=logging.INFO,format='%(asctime)s %(levelname)s %(message)s')
log=logging.getLogger('review')

class Result(BaseModel):
    """固定型別讓 API、ETL 與報告共用相同契約；缺值保留 null。"""
    product: str|None=None
    product_evidence: list[dict]=Field(default_factory=list)
    batch: str|None=None
    expiry: str|None=None
    expiry_raw: str|None=None
    serial_number: str|None=None
    field_candidates: dict[str,list[str]]=Field(default_factory=dict)
    ocr_attempts: list[dict]=Field(default_factory=list)
    full_text: str
    confidence: float|None=Field(default=None,ge=0,le=100)
    mode: Literal['text_demo','ocr']='text_demo'
    review_required: Literal[True]=True
    warnings: list[str]=Field(default_factory=list)

    @field_validator('expiry')
    @classmethod
    def valid_expiry(cls,value):
        if value is not None:
            if not re.fullmatch(r'\d{4}-\d{2}-\d{2}',value): raise ValueError('效期必須為 ISO 日期')
            date.fromisoformat(value)
        return value

class TextInput(BaseModel):
    text: str=Field(min_length=1,max_length=100000)

def normalize_date(raw):
    """接受年/月/日、年/月與月/年；不猜測 O/0；年月依明示規則取月底。"""
    raw=re.sub(r'\s+','',raw.strip())
    if re.fullmatch(r'\d{2}/\d{4}',raw):
        mo,y=raw.split('/');raw=y+'/'+mo
    m=re.fullmatch(r'(\d{4})[-/](\d{2})(?:[-/](\d{2}))?',raw)
    if not m:return None
    try:
        y,mo=int(m[1]),int(m[2]);d=int(m[3]) if m[3] else calendar.monthrange(y,mo)[1]
        return date(y,mo,d).isoformat()
    except ValueError:return None

def parse(text,confidence=None,mode='text_demo'):
    """允許標籤後有冒號或空格；值原樣保留，絕不替換形近字。"""
    def field(labels,pattern):
        m=re.search(r'(?im)^\s*(?:'+labels+r')(?:[ \t]*[:：][ \t]*|[ \t]+)('+pattern+r')',text)
        return m[1].strip() if m else None
    product=field('Product|品名',r'[^\n\r]+')
    batch=field('LOT|Batch|批號',r'[A-Za-z0-9][A-Za-z0-9-]*')
    raw=field('EXP|Expiry|效期',r'[^\n\r]+')
    # 日期後常有 OCR 雜訊，只接受完整日期字串的第一段，保留原始整行作稽核。
    match=re.match(r'(\d{4}[-/]\d{2}(?:[-/]\d{2})?|\d{2}/\d{4})(?![\d/])',raw or '')
    expiry=normalize_date(match[0]) if match else None
    serial=field('SN|S/N|Serial(?: number)?|序號',r'[A-Za-z0-9]+(?:[ \t]+[0-9]+)*')
    if serial:serial=re.sub(r'[ \t]+','',serial)
    warnings=[]
    for name,value in [('品名',product),('批號',batch),('效期',expiry)]:
        if value is None:warnings.append(name+'缺值或格式無法確認，請對照原圖')
    if match and len(re.findall(r'[-/]',match[0]))==1:warnings.append('年月效期依示範規則換算月底，須由 SOP 確認；原印字保留於 expiry_raw')
    if confidence is not None and confidence<85:warnings.append('OCR 平均信心低於 85，優先人工覆核')
    if mode=='text_demo':warnings.append('文字備援模式，未執行 OCR')
    return Result(product=product,batch=batch,expiry=expiry,expiry_raw=raw,serial_number=serial,full_text=text,confidence=confidence,mode=mode,warnings=warnings)


def pages_for(source,folder):
    """保留原檔；PDF 轉成逐頁影像，最多十頁以限制資源。"""
    if source.suffix.lower()=='.pdf':
        from pdf2image import convert_from_path, pdfinfo_from_path
        if pdfinfo_from_path(source)['Pages']>10: raise ValueError('PDF 最多十頁')
        imgs=convert_from_path(source,dpi=150,first_page=1,last_page=10)
    else:
        with Image.open(source) as im:
            if im.width*im.height>25000000: raise ValueError('圖片超過 2500 萬像素')
            imgs=[ImageOps.exif_transpose(im).convert('RGB')]
    paths=[]
    for i,im in enumerate(imgs):
        p=folder/f'page-{i+1}.png'; im.save(p); paths.append(p)
    return paths

def ocr_pages(paths,lang='eng'):
    from ocr_upgrade import attempts_for
    attempts=[]
    for page,p in enumerate(paths,1):
        with Image.open(p) as im:
            for attempt in attempts_for(im.convert('RGB'),lang):
                attempt['page']=page;attempts.append(attempt)
    base=[a for a in attempts if a['source']=='original']
    full_text='\n\n--- PAGE ---\n\n'.join(a['text'] for a in base)
    r=parse(full_text,sum(a['confidence'] for a in base)/len(base) if base else 0,'ocr')
    r.ocr_attempts=attempts
    # 多次辨識不投機補值；不同候選即留空，讓覆核者對照原圖選定。
    candidates={name:[] for name in ['batch','expiry','serial_number']}
    for attempt in attempts:
        extracted=parse(attempt['text'],attempt['confidence'],'ocr')
        for name in candidates:
            value=getattr(extracted,name)
            # 過短的增強辨識片段視為雜訊；原圖明確解析仍保留。
            if attempt['source']!='original' and value and name in ['batch','serial_number']:
                if len(value)<(2 if name=='batch' else 6):value=None
            if value and value not in candidates[name]:candidates[name].append(value)
        if extracted.expiry and (not normalize_date(r.expiry_raw or '') or not r.expiry_raw):r.expiry_raw=extracted.expiry_raw
    from product_parser import product_from_attempts
    product_candidates,evidence=product_from_attempts(attempts)
    r.product_evidence=evidence
    if not r.product and product_candidates:
        r.product=product_candidates[0] if len(product_candidates)==1 else None
        if r.product:
            r.warnings=[w for w in r.warnings if not w.startswith('品名缺值')]
            r.warnings.append('品名由同頁 OCR 品牌、規格與劑型組合，須逐字人工確認；來源保留於 product_evidence')
        else:r.warnings.append('品名有不同候選，未自動選定，請人工確認')
    if product_candidates:candidates['product']=product_candidates
    r.field_candidates=candidates
    for name,values in candidates.items():
        setattr(r,name,values[0] if len(values)==1 else None)
        if len(values)>1:r.warnings.append(name+'有不同 OCR 候選：'+' / '.join(values)+'；未自動選定，請人工確認')
    for name,field in [('批號','batch'),('效期','expiry')]:
        if getattr(r,field):r.warnings=[w for w in r.warnings if not w.startswith(name+'缺值')]
    if r.expiry and normalize_date(r.expiry_raw or '') and not any('換算月底' in w for w in r.warnings):
        if len(re.findall(r'[-/]',r.expiry_raw))==1:r.warnings.append('年月效期依示範規則換算月底，須由 SOP 確認；非自動放行判定')
    if len(attempts)>len(base):r.warnings.append('已嘗試小字區域放大與銳化；辨識來源、座標、原文與分數保留於 ocr_attempts')
    return r


def report(r,images,target):
    """原圖與辨識文字分頁展示；人工檢查與簽核欄永遠留白。"""
    d=Document(); sec=d.sections[0]
    sec.top_margin=sec.bottom_margin=Inches(.7)
    style=d.styles['Normal']; style.font.name='Arial'; style.font.size=Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'),'Heiti TC')
    for n in ['Title','Heading 1','Heading 2']:
        d.styles[n].font.color.rgb=RGBColor(0,0,0)
        d.styles[n].element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Heiti TC')
        for border in list(d.styles[n].element.findall('.//'+qn('w:pBdr'))): border.getparent().remove(border)
    d.add_heading('進口醫藥產品批次文件審核輔助報告',0)
    d.add_paragraph('文件辨識輔助紀錄。本報告僅整理 OCR 或輸入結果，所有欄位均待人工覆核，不構成產品放行決定或法規合規驗證。')
    d.add_heading('擷取結果',1)
    for label,v in [('品名',r.product),('批號',r.batch),('原印效期',r.expiry_raw),('換算效期',r.expiry),('序號 SN',r.serial_number),('模式',r.mode),('平均信心',r.confidence)]:
        d.add_paragraph(f'{label}：{v if v is not None else "未辨識"}')
    for w in r.warnings: d.add_paragraph('提示：'+w)
    for field,values in r.field_candidates.items():
        if len(values)>1:d.add_paragraph('待覆核候選 '+field+'：'+' / '.join(values))
    d.add_heading('人工覆核',1)
    d.add_paragraph('□ 品名及規格  □ 批號  □ 效期與 SOP  □ 仿單完整性\n□ 字體與版面  □ 色塊與印刷  □ 包裝與核准樣本一致性\n色塊與版面僅提示人工對照；系統未執行色差或版面合格判定。\n差異與處置：________________________________\n覆核者：________________  日期：________________\n品質單位依既有程序另行決定處置：________________')
    for i,p in enumerate(images):
        d.add_page_break(); d.add_heading(f'原圖 第 {i+1} 頁',1); d.add_picture(str(p),width=Inches(6))
    d.add_page_break(); d.add_heading('完整辨識文字',1)
    d.add_paragraph(r.full_text or '無辨識文字，不得補造內容。')
    for paragraph in d.paragraphs:
        for border in list(paragraph._p.findall('.//'+qn('w:pBdr'))): border.getparent().remove(border)
    d.save(target)

def save_etl(records,db):
    """以原始文字 SHA256 去重，SQL 參數綁定避免拼接資料。"""
    inserted=0
    with sqlite3.connect(db) as c:
        c.execute('CREATE TABLE IF NOT EXISTS batches (id TEXT PRIMARY KEY, product TEXT, batch TEXT, expiry TEXT, payload TEXT NOT NULL)')
        for r in records:
            key=hashlib.sha256(r.full_text.encode()).hexdigest()
            cur=c.execute('INSERT OR IGNORE INTO batches VALUES (?,?,?,?,?)',(key,r.product,r.batch,r.expiry,r.model_dump_json()))
            inserted+=cur.rowcount
        total,missing=c.execute('SELECT COUNT(*), COALESCE(SUM(expiry IS NULL),0) FROM batches').fetchone()
    return dict(inserted=inserted,duplicates=len(records)-inserted,total=total,missing_expiry=missing)

from rag import Query, Answer, retrieve

app=FastAPI(title='批次文件審核輔助 MVP',version='1.0.0')
@app.post('/rag',response_model=Answer)
def rag_api(body:Query):
    from gemini_rag import answer_question
    return answer_question(body.question, body.use_llm)

@app.get('/health')
def health(): return {'status':'ok','tesseract':bool(shutil.which('tesseract')),'poppler':bool(shutil.which('pdftoppm'))}
@app.get('/',response_class=HTMLResponse)
def home():
    # 畫面獨立存放，讓流程邏輯與視覺設計容易維護。
    return (ROOT/'web'/'index.html').read_text(encoding='utf-8')

def complete(r,paths,folder):
    report(r,paths,folder/'report.docx'); (folder/'result.json').write_text(r.model_dump_json(indent=2),encoding='utf-8')
    log.info('report_created job=%s mode=%s',folder.name,r.mode)
    return {'result':r.model_dump(),'report_url':f'/reports/{folder.name}'}
@app.post('/text')
def text_api(body:TextInput):
    folder=OUT/uuid.uuid4().hex; folder.mkdir()
    return complete(parse(body.text),[],folder)
@app.post('/ocr')
async def ocr_api(file:UploadFile=File(...)):
    suffix=Path(file.filename or '').suffix.lower()
    if suffix not in {'.png','.jpg','.jpeg','.pdf'}: raise HTTPException(415,'僅支援 PNG JPG PDF')
    content=await file.read(10*1024*1024+1)
    if len(content)>10*1024*1024: raise HTTPException(413,'檔案上限 10 MB')
    folder=OUT/uuid.uuid4().hex; folder.mkdir(); source=folder/('original'+suffix); source.write_bytes(content)
    try:
        paths=pages_for(source,folder)
        import os
        r=ocr_pages(paths,os.getenv('OCR_LANG','eng'))
        return complete(r,paths,folder)
    except Exception as e:
        log.exception('ocr_failed job=%s',folder.name)
        raise HTTPException(422,'OCR 失敗；請檢查檔案、Tesseract 語言包及 Poppler，或使用文字備援。') from e
@app.get('/reports/{job}')
def download(job:str):
    if not re.fullmatch('[0-9a-f]{32}',job): raise HTTPException(404)
    p=OUT/job/'report.docx'
    if not p.exists(): raise HTTPException(404)
    return FileResponse(p,filename='review-report.docx')

if __name__=='__main__':
    a=argparse.ArgumentParser(); a.add_argument('--text',type=Path); a.add_argument('--file',type=Path); a.add_argument('--lang',default='eng'); args=a.parse_args()
    folder=OUT/uuid.uuid4().hex; folder.mkdir()
    if args.file:
        source=folder/('original'+args.file.suffix.lower()); shutil.copy2(args.file,source)
        paths=pages_for(source,folder); r=ocr_pages(paths,args.lang)
    elif args.text: paths=[]; r=parse(args.text.read_text(encoding='utf-8'))
    else: a.error('請指定 --text 或 --file')
    print(json.dumps(complete(r,paths,folder),ensure_ascii=False,indent=2))

@app.get("/rag/status")
def rag_status():
    from gemini_rag import configuration
    return configuration()
